from __future__ import annotations

import asyncio
import io
import logging
import os
import re
import shutil
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from core.settings import get_settings
from services.database import get_db
from services.llm_factory import get_llm_gateway
from core.skill_engine.base import SkillContext

logger = logging.getLogger(__name__)

from services.middleware.api_key import require_any_auth

router = APIRouter(dependencies=[Depends(require_any_auth)])

UPLOAD_DIR = Path("./uploads/formatted")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

TEMPLATE_DIR = Path("./templates")
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)

EXPORT_DIR = Path("./uploads/exports")
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

MAX_UPLOAD_BYTES = 50 * 1024 * 1024


def _save_upload(file: UploadFile, prefix: str = "") -> Path:
    if file.size and file.size > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    name = Path(file.filename or "upload").name
    safe = "".join(c for c in name if c.isalnum() or c in ("-", "_", "."))
    dest = UPLOAD_DIR / f"{prefix}{safe}"
    content = file.file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    dest.write_bytes(content)
    return dest


@contextmanager
def _temp_upload(path: Path):
    """请求级临时文件清理: 只删本次请求 _save_upload 生成的临时文件。

    请求正常结束或抛异常都会清理; 不碰 uploads 目录里的历史文件,
    也不删 skill 产出的导出产物(用户要用 /download 取回)。
    """
    try:
        yield path
    finally:
        try:
            if path is not None:
                Path(path).unlink(missing_ok=True)
        except OSError:
            logger.warning("临时上传文件清理失败: %s", path)


_SAFE_NAME_RE = re.compile(r"^[A-Za-z0-9_.\-]+$")


def _safe_template_name(name: str) -> str:
    """模板名白名单清洗: 只取基名, 仅允许字母数字/下划线/连字符/点, 防路径穿越。"""
    base = Path((name or "").replace("\\", "/")).name
    if not base or not _SAFE_NAME_RE.match(base) or base.startswith("."):
        raise HTTPException(
            status_code=400,
            detail="非法模板名: 仅允许字母、数字、下划线、连字符和点",
        )
    return base


@router.post("/format")
async def format_document(
    file: UploadFile = File(...),
    template: str = "default",
    mode: str = Query("format", description="format|check|diff|beautify"),
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file)
    with _temp_upload(temp_path):
        from services.format.skills.docx_format_skill import DocxFormatSkill

        gateway = get_llm_gateway()
        skill = DocxFormatSkill()
        ctx = SkillContext(
            project_id="",
            db=db,
            llm=gateway,
            parameters={
                "file_path": str(temp_path),
                "template": template,
                "mode": mode,
            },
        )
        skill_result = await skill.safe_execute(ctx)

        if skill_result.success and skill_result.data:
            result_data = skill_result.data
            output_path = result_data.get("output_path", "")
            if output_path and Path(output_path).exists():
                result_data["file_name"] = Path(output_path).name
                try:
                    result_data["file_size"] = Path(output_path).stat().st_size
                except OSError:
                    pass

            return {"success": True, "data": result_data}

        return {"success": False, "error": skill_result.error}


@router.post("/check-format")
async def check_format(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="check_")
    with _temp_upload(temp_path):
        from services.format.skills.docx_format_skill import DocxFormatSkill

        gateway = get_llm_gateway()
        skill = DocxFormatSkill()
        ctx = SkillContext(
            project_id="",
            db=db,
            llm=gateway,
            parameters={
                "file_path": str(temp_path),
                "template": template,
                "mode": "check",
            },
        )
        skill_result = await skill.safe_execute(ctx)

        if skill_result.success:
            return {"success": True, "data": skill_result.data}
        return {"success": False, "error": skill_result.error}


@router.post("/diff-format")
async def diff_format(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="diff_")
    with _temp_upload(temp_path):
        from services.format.skills.docx_format_skill import DocxFormatSkill

        gateway = get_llm_gateway()
        skill = DocxFormatSkill()
        ctx = SkillContext(
            project_id="",
            db=db,
            llm=gateway,
            parameters={
                "file_path": str(temp_path),
                "template": template,
                "mode": "diff",
            },
        )
        skill_result = await skill.safe_execute(ctx)

        if skill_result.success:
            return {"success": True, "data": skill_result.data}
        return {"success": False, "error": skill_result.error}


@router.post("/beautify")
async def beautify_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="beautify_")
    with _temp_upload(temp_path):
        from services.format.skills.docx_format_skill import DocxFormatSkill

        gateway = get_llm_gateway()
        skill = DocxFormatSkill()
        ctx = SkillContext(
            project_id="",
            db=db,
            llm=gateway,
            parameters={
                "file_path": str(temp_path),
                "mode": "beautify",
            },
        )
        skill_result = await skill.safe_execute(ctx)

        if skill_result.success and skill_result.data:
            return {"success": True, "data": skill_result.data}
        return {"success": False, "error": skill_result.error}


def _q(name: str) -> str:
    """对下载文件名做百分号编码(Starlette 用 latin-1 编码响应头,中文会 500)。"""
    from urllib.parse import quote

    return quote(name)


@router.post("/export-doc")
async def export_to_doc(
    file: UploadFile = File(...),
    template: str = "default",
    apply_format: bool = Form(True),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版(可选) → 转 doc 字节流下载。"""
    src = _save_upload(file, prefix="srcdocx_")
    with _temp_upload(src):
        docx_to_send: Path = src
        if apply_format:
            formatted = await _run_format_skill(src, template, db)
            if formatted:
                docx_to_send = formatted

        from services.format.skills.doc_export_skill import DocExportSkill
        skill = DocExportSkill()
        if not docx_to_send.exists():
            raise HTTPException(status_code=500, detail="排版产物丢失")

        # _try_libreoffice_doc 内部是 subprocess.run(超时 120s) 的同步阻塞调用,
        # 放到线程池执行, 避免卡住事件循环。
        result = await asyncio.to_thread(
            skill._try_libreoffice_doc, docx_to_send, docx_to_send.parent
        )
        if result.get("success"):
            doc_path = Path(result["output_path"])
            data = doc_path.read_bytes()
            try:
                doc_path.unlink(missing_ok=True)
            except OSError:
                pass
        else:
            # 本服务器未安装 LibreOffice(约400MB,1.8G内存的机器装不动)。
            # 回退: 直接下发 docx 字节流,扩展名仍用 .doc,
            # Word/WPS 会按文件真实格式正常打开。
            data = docx_to_send.read_bytes()

        if docx_to_send != src:
            try:
                docx_to_send.unlink(missing_ok=True)
            except OSError:
                pass

        filename = f"{Path(file.filename or 'document').stem}.doc"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/msword",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_q(filename)}"},
        )


def _docx_to_pdf_weasyprint(docx_path: Path) -> bytes | None:
    """docx → mammoth(HTML) → weasyprint(PDF) 兜底通路。

    容器里装了 mammoth + weasyprint,但没装 LibreOffice/docx2pdf/reportlab,
    所以这是本机唯一可用的 PDF 转换方式(依赖系统已安装中文字体)。
    """
    try:
        import mammoth
        from weasyprint import HTML
    except Exception:
        return None
    try:
        with open(docx_path, "rb") as fh:
            body = mammoth.convert_to_html(fh).value or ""
        if not body.strip():
            return None
        html = (
            "<html><head><meta charset='utf-8'><style>"
            "@page { size: A4; margin: 2.54cm 3.17cm; }"
            "body { font-family: 'WenQuanYi Zen Hei','Noto Sans CJK SC','SimSun',sans-serif;"
            " font-size: 12pt; line-height: 1.65; }"
            "h1 { font-size: 18pt; } h2 { font-size: 15pt; } h3 { font-size: 13pt; }"
            "table { border-collapse: collapse; width: 100%; }"
            "td, th { border: 1px solid #666; padding: 4px 6px; font-size: 10.5pt; }"
            "</style></head><body>" + body + "</body></html>"
        )
        return HTML(string=html).write_pdf()
    except Exception:
        return None


@router.post("/export-pdf")
async def export_to_pdf(
    file: UploadFile = File(...),
    template: str = "default",
    apply_format: bool = Form(True),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版(可选) → 转 pdf 字节流下载。"""
    src = _save_upload(file, prefix="srcdocx_")
    with _temp_upload(src):
        docx_to_send: Path = src
        if apply_format:
            formatted = await _run_format_skill(src, template, db)
            if formatted:
                docx_to_send = formatted

        from services.format.skills.pdf_export_skill import PdfExportSkill
        gateway = get_llm_gateway()
        skill = PdfExportSkill()
        out_dir = docx_to_send.parent
        ctx = SkillContext(
            project_id="",
            db=db,
            llm=gateway,
            parameters={
                "input_path": str(docx_to_send),
                "output_dir": str(out_dir),
            },
        )
        # PdfExportSkill.safe_execute 是 async, 其内部 subprocess.run 阻塞属于该
        # skill 自身实现, 按约定不在此处再包线程池(见报告说明)。
        skill_result = await skill.safe_execute(ctx)

        data = None
        if skill_result.success and skill_result.data and skill_result.data.get("output_path"):
            pdf_path = Path(skill_result.data["output_path"])
            if pdf_path.exists():
                data = pdf_path.read_bytes()
                try:
                    pdf_path.unlink(missing_ok=True)
                except OSError:
                    pass
        if not data:
            # 技能内三条策略(LibreOffice/docx2pdf/reportlab)在本机都没装,
            # 用 mammoth + weasyprint 兜底; weasyprint 渲染是 CPU 阻塞,
            # 同样放线程池执行。
            data = await asyncio.to_thread(_docx_to_pdf_weasyprint, docx_to_send)
        if not data:
            if docx_to_send != src:
                try:
                    docx_to_send.unlink(missing_ok=True)
                except OSError:
                    pass
            raise HTTPException(
                status_code=500,
                detail=(skill_result.error or "PDF 转换失败")
                + "；备用方案 mammoth+weasyprint 也失败(请确认系统已安装中文字体)",
            )

        if docx_to_send != src:
            try:
                docx_to_send.unlink(missing_ok=True)
            except OSError:
                pass

        filename = f"{Path(file.filename or 'document').stem}.pdf"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_q(filename)}"},
        )


@router.post("/export-formatted-docx")
async def export_formatted_docx(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版 → 直接下载 docx 字节流。"""
    src = _save_upload(file, prefix="srcdocx_")
    with _temp_upload(src):
        formatted = await _run_format_skill(src, template, db) or src
        data = formatted.read_bytes()
        try:
            if formatted != src:
                formatted.unlink(missing_ok=True)
        except OSError:
            pass
        filename = f"{Path(file.filename or 'document').stem}_formatted.docx"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_q(filename)}"},
        )


async def _run_format_skill(src: Path, template: str, db: AsyncSession) -> Path | None:
    from services.format.skills.docx_format_skill import DocxFormatSkill
    gateway = get_llm_gateway()
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(src),
            "template": template,
            "mode": "format",
        },
    )
    result = await skill.safe_execute(ctx)
    if not result.success or not result.data:
        logger.warning("format skill 失败: %s", result.error)
        return None
    out = result.data.get("output_path", "")
    if out and Path(out).exists():
        return Path(out)
    return None


@router.post("/from-project/{project_id}")
async def format_from_project(
    project_id: str,
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    """从项目章节组装为 docx 文件（无需上传）"""
    from sqlalchemy import select
    from services.models import Project, Chapter

    proj_result = await db.execute(select(Project).where(Project.id == project_id))
    project = proj_result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    ch_result = await db.execute(
        select(Chapter).where(Chapter.project_id == project_id).order_by(Chapter.sort_order)
    )
    chapters = ch_result.scalars().all()
    if not chapters:
        raise HTTPException(
            status_code=400,
            detail="项目尚无任何章节大纲，请先到「投标生成」完成大纲生成。",
        )

    empty_chapters = [ch for ch in chapters if not (ch.content and ch.content.strip())]
    if empty_chapters:
        empty_titles = [ch.title for ch in empty_chapters[:5]]
        detail_msg = (
            f"项目存在 {len(empty_chapters)}/{len(chapters)} 个章节尚未生成正文，"
            f"无法执行文档输出。请先到「投标生成」完成正文生成。"
            f"未生成章节示例: {', '.join(empty_titles)}"
        )
        if len(empty_chapters) == len(chapters):
            raise HTTPException(status_code=400, detail=detail_msg)
        if len(empty_chapters) / len(chapters) > 0.5:
            raise HTTPException(
                status_code=400,
                detail=detail_msg + "（空内容章节超过50%，已禁止操作）",
            )

    title = f"{project.name} - 投标文件"

    import tempfile
    from docx import Document
    from docx.shared import Pt
    from pathlib import Path
    from core.skill_engine.base import SkillContext
    from services.format.skills.docx_format_skill import DocxFormatSkill

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    h1 = doc.add_heading(title, level=0)

    for ch in chapters:
        if not (ch.content and ch.content.strip()):
            continue
        doc.add_heading(ch.title, level=1)
        for line in ch.content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            else:
                doc.add_paragraph(line)

    output_dir = Path(tempfile.gettempdir()) / "bidmaster_format"
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[^\w\u4e00-\u9fff\-_]', '_', project.name)[:60] or "bid"
    src_path = output_dir / f"{safe_name}_src_{project_id[:8]}.docx"
    doc.save(str(src_path))

    from services.llm_factory import get_llm_gateway
    gateway = get_llm_gateway()
    skill = DocxFormatSkill()

    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={"file_path": str(src_path), "template": template, "mode": "format"},
    )
    result = await skill.safe_execute(ctx)
    if not result.success:
        try:
            src_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=result.error or "项目章节组装失败")

    final_output = result.data.get("output_path") if isinstance(result.data, dict) else None

    return {
        "success": True,
        "output_path": final_output,
        "project_name": project.name,
        "chapter_count": len(chapters),
        "generated_chapter_count": len(chapters) - len(empty_chapters),
        "empty_chapter_count": len(empty_chapters),
        "chapters": [
            {
                "id": str(c.id),
                "title": c.title,
                "word_count": c.word_count or 0,
                "status": c.status,
                "has_content": bool(c.content and c.content.strip()),
            }
            for c in chapters
        ],
    }


@router.get("/templates")
async def list_templates():
    templates = [
        {"name": "default", "description": "默认标书排版模板(仿宋正文+黑体标题)"},
        {"name": "government", "description": "政府采购标书模板(方正小标宋+仿宋)"},
        {"name": "engineering", "description": "工程标书模板(宋体正文+黑体标题)"},
    ]

    for tpl_file in TEMPLATE_DIR.glob("*.yaml"):
        name = tpl_file.stem
        if not any(t["name"] == name for t in templates):
            templates.append({"name": name, "description": f"自定义模板: {name}"})

    return {"templates": templates}


@router.get("/templates/{template_name}")
async def get_template_detail(template_name: str):
    from services.format.skills.docx_format_skill import DEFAULT_FORMAT_CONFIG

    safe_name = _safe_template_name(template_name)
    template_path = TEMPLATE_DIR / f"{safe_name}.yaml"
    if template_path.exists():
        try:
            import yaml
            with open(template_path, encoding="utf-8") as f:
                custom_config = yaml.safe_load(f) or {}
            merged = {**DEFAULT_FORMAT_CONFIG, **custom_config}
        except Exception:
            merged = DEFAULT_FORMAT_CONFIG
    else:
        merged = DEFAULT_FORMAT_CONFIG

    return {"name": safe_name, "config": merged}


@router.put("/templates/{template_name}")
async def save_template(template_name: str, config: dict):
    import yaml

    safe_name = _safe_template_name(template_name)
    template_path = TEMPLATE_DIR / f"{safe_name}.yaml"
    with open(template_path, "w", encoding="utf-8") as f:
        yaml.dump(config, f, allow_unicode=True, default_flow_style=False)

    return {"success": True, "name": safe_name}


@router.delete("/templates/{template_name}")
async def delete_template(template_name: str):
    safe_name = _safe_template_name(template_name)
    if safe_name in ("default", "government", "engineering"):
        return {"success": False, "error": "内置模板不可删除"}

    template_path = TEMPLATE_DIR / f"{safe_name}.yaml"
    if template_path.exists():
        template_path.unlink()
        return {"success": True}
    return {"success": False, "error": "模板不存在"}


# 按扩展名返回正确的 Content-Type(原先无论什么类型都返回 docx 的 MIME)
_DOWNLOAD_MEDIA_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".pdf": "application/pdf",
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/markdown; charset=utf-8",
    ".html": "text/html; charset=utf-8",
    ".zip": "application/zip",
}


@router.get("/download")
async def download_formatted_file(path: str):
    """下载由 /format 产生的产物（限定 uploads/formatted 目录）。"""
    base = Path(os.path.realpath(UPLOAD_DIR))
    candidate = Path(path)
    if not candidate.is_absolute():
        candidate = base / path
    # 先 realpath 解析(吃掉 ../ 与符号链接), 再用 commonpath 严格判断是否在目录内;
    # 不用 startswith, 否则 uploads/formatted_evil 这类同前缀目录会绕过。
    p = Path(os.path.realpath(candidate))
    try:
        inside = os.path.commonpath([str(base), str(p)]) == str(base) and p != base
    except ValueError:
        # 不同盘符/无法计算公共路径
        inside = False
    if not inside:
        raise HTTPException(status_code=400, detail="非法路径")
    if not p.exists() or not p.is_file():
        raise HTTPException(status_code=404, detail="文件不存在")
    media_type = _DOWNLOAD_MEDIA_TYPES.get(p.suffix.lower(), "application/octet-stream")
    return StreamingResponse(
        p.open("rb"),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_q(p.name)}"},
    )
