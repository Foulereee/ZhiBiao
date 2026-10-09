# -*- coding: utf-8 -*-
"""字体归一化实测脚本（可复现 README「修复效果实测」一节的数据）。

做三件事：
  1. 用**有问题的写法**（只设 run.font.name，不写 w:eastAsia）生成一份样例 docx；
  2. 调用 services/format/font_utils.py 的 normalize_docx_fonts() 就地归一化；
  3. 打印归一化前后 w:rFonts 的字体取值与 w:eastAsia 缺失数。

用法：
    python tools/font_normalize_evidence.py [输出目录]
    # 不传输出目录则写到系统临时目录，并打印路径

为什么需要它：Word 对中文字符取 w:eastAsia，而 python-docx 的 run.font.name 只写
w:ascii / w:hAnsi —— 缺失时中文回退 docDefaults 的主题字体（等线/Calibri），
数字与英文仍按 ascii 用宋体，于是同一行出现两种字体。本脚本把这个过程量化出来。
"""
from __future__ import annotations

import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from docx import Document                              # noqa: E402
from docx.shared import Pt                             # noqa: E402

from services.format.font_utils import normalize_docx_fonts  # noqa: E402

_RFONTS_RE = re.compile(r"<w:rFonts\b[^>]*/?>")
_DOCDEFAULTS_RE = re.compile(r"<w:docDefaults>.*?</w:docDefaults>", re.S)


def build_sample(path: str) -> None:
    """按"只写 ascii/hAnsi"的有问题写法生成样例文档。"""
    doc = Document()
    head = doc.add_heading("投标文件（字体问题复现）", level=1)
    for run in head.runs:
        run.font.name = "宋体"

    p1 = doc.add_paragraph()
    r1 = p1.add_run("投标报价为 1,234,567.89 元")
    r1.font.name = "宋体"
    r1.font.size = Pt(12)

    p2 = doc.add_paragraph()
    p2.add_run("Total Price: 1,234,567.89 CNY（含税）").font.name = "宋体"

    p3 = doc.add_paragraph()
    p3.add_run("排版链路残留：").font.name = "仿宋_GB2312"
    p3.add_run("本段其余文字仍为宋体").font.name = "宋体"

    doc.save(path)


def measure(path: str) -> dict:
    """统计 document.xml / styles.xml 里的字体分布。"""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        styles = z.read("word/styles.xml").decode("utf-8") if "word/styles.xml" in z.namelist() else ""

    rfonts = _RFONTS_RE.findall(xml)
    dd = _DOCDEFAULTS_RE.search(styles)
    dd_ea: list[str] = []
    if dd:
        for tag in _RFONTS_RE.findall(dd.group(0)):
            dd_ea += re.findall(r'w:eastAsia="([^"]+)"', tag)

    return {
        "rFonts 总数": len(rfonts),
        "缺 w:eastAsia 的 rFonts": len([t for t in rfonts if "w:eastAsia=" not in t]),
        "eastAsia 取值": sorted({m for t in rfonts for m in re.findall(r'w:eastAsia="([^"]+)"', t)}),
        "ascii 取值": sorted({m for t in rfonts for m in re.findall(r'w:ascii="([^"]+)"', t)}),
        "主题引用(Theme)取值": sorted({m for t in rfonts
                                    for m in re.findall(r'w:(?:ascii|hAnsi|eastAsia|c)Theme="([^"]+)"', t)}),
        "docDefaults.eastAsia": dd_ea,
    }


def main() -> int:
    out = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="font-evidence-")
    os.makedirs(out, exist_ok=True)
    before = os.path.join(out, "01_修复前.docx")
    after = os.path.join(out, "02_修复后.docx")

    build_sample(before)
    m_before = measure(before)

    shutil.copyfile(before, after)
    changed = normalize_docx_fonts(after)
    m_after = measure(after)

    print("输出目录: %s\n" % out)
    print("归一化改写的部件: %s\n" % (changed or "（无）"))
    width = max(len(k) for k in m_before)
    def _fmt(value: object) -> str:
        # 注意：0 / False 是有效观测值，不能当"空"处理
        if value is None or value == "" or value == []:
            return "（空）"
        return str(value)

    print("%-*s | %-34s | %s" % (width, "观测项", "修复前", "修复后"))
    print("%s-+-%s-+-%s" % ("-" * width, "-" * 34, "-" * 34))
    for key in m_before:
        print("%-*s | %-34s | %s" % (width, key, _fmt(m_before[key]), _fmt(m_after.get(key))))

    with open(os.path.join(out, "测量结果.json"), "w", encoding="utf-8") as f:
        json.dump({"before": m_before, "after": m_after, "changed_parts": changed}, f,
                  ensure_ascii=False, indent=2)
    print("\n缺少 w:eastAsia 的 rFonts: %d → %d" %
          (m_before["缺 w:eastAsia 的 rFonts"], m_after["缺 w:eastAsia 的 rFonts"]))
    ok = m_after["缺 w:eastAsia 的 rFonts"] == 0 and bool(m_after["eastAsia 取值"])
    print("结论: %s" % ("归一化生效 ✅" if ok else "归一化未生效 ❌"))
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())