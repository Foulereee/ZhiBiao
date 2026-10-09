from __future__ import annotations

import logging

from core.doc_engine.parsers.base import ParsedDocument

logger = logging.getLogger(__name__)


def _iter_block_items(doc):
    """按文档真实顺序迭代「段落 / 表格」。

    python-docx 的 doc.paragraphs 与 doc.tables 是两个互不相干的列表，
    直接分别遍历会丢掉表格在正文中的位置。而招标文件里最关键的内容
    （评分标准表、投标人须知前附表、资格要求表、报价表）几乎都在表格里，
    只取 paragraphs 会导致这些表格整表丢失，下游「评分矩阵 → 大纲对齐 →
    正文生成 → 评分覆盖」全部拿不到真实评分项。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _table_to_text(table, index: int) -> tuple[str, dict]:
    """把一张表转成可读文本（喂给 LLM）以及结构化记录（存 tables）。"""
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])

    headers = rows[0] if rows else []
    body_rows = rows[1:] if len(rows) > 1 else []

    lines = [f"【表{index}】"]
    for r in rows:
        cells = [c.replace("\n", " ").strip() for c in r]
        lines.append(" | ".join(cells))
    text = "\n".join(lines)

    return text, {
        "index": index,
        "headers": headers,
        "rows": body_rows,
    }


class DocxParser:
    def parse(self, file_path: str) -> ParsedDocument:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(file_path)
        text_parts: list[str] = []
        tables: list[dict] = []
        para_count = 0
        table_idx = 0

        try:
            blocks = list(_iter_block_items(doc))
        except Exception as e:  # 极端畸形文档：退回旧行为，至少不丢段落
            logger.warning(f"[DocxParser] 按文档顺序遍历失败，退回段落优先模式: {e}")
            blocks = list(doc.paragraphs)

        for block in blocks:
            if isinstance(block, Paragraph):
                para_count += 1
                text = block.text.strip()
                if text:
                    text_parts.append(text)
            elif isinstance(block, Table):
                table_text, record = _table_to_text(block, table_idx)
                tables.append(record)
                text_parts.append(table_text)
                table_idx += 1

        return ParsedDocument(
            text="\n".join(text_parts),
            tables=tables,
            images=[],
            metadata={
                "paragraph_count": para_count,
                "table_count": len(tables),
                "parser": "python-docx",
                # 标记表格已并入 text，下游不必再次合并（避免重复）
                "tables_inlined": True,
            },
        )